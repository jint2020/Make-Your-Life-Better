"""转换引擎的单元测试：直接调 engine.convert_document，不经过 HTTP / 子进程。"""

from __future__ import annotations

import pytest

from app import engine, postprocess
from app.ole2 import sniff_ole2

from .fixtures import (
    make_chinese_pdf_with_table,
    make_csv_gbk,
    make_csv_many_rows,
    make_csv_utf8,
    make_docx_image_only,
    make_docx_with_table_and_image,
    make_encrypted_pdf,
    make_epub,
    make_html_with_images,
    make_ole2_encrypted,
    make_ole2_legacy,
    make_pdf_image_only,
    make_plain_zip,
    make_pptx_image_only,
    make_pptx_with_picture,
    make_xlsx_bogus_dimension,
    make_xlsx_many_rows,
    make_xlsx_single_column,
    make_xlsx_with_dates,
)


def test_docx_chinese_text_and_table() -> None:
    result = engine.convert_document(make_docx_with_table_and_image(), ".docx", 50_000)
    md = result["markdown"]
    assert "文档标题" in md
    assert "张三" in md
    assert "100" in md
    # 表格列存在（| 姓名 | 分数 | 或者两边各有一个空格）
    assert "姓名" in md and "分数" in md


def test_docx_image_becomes_placeholder() -> None:
    result = engine.convert_document(make_docx_with_table_and_image(), ".docx", 50_000)
    md = result["markdown"]
    # 内嵌图片 → 占位符；不应该有 data:image/...
    assert "[图片]" in md
    assert "data:image" not in md
    # 没有 alt 文本时就是纯 [图片]，不要多出中文
    assert "[图片：" not in md


def test_docx_image_only_no_text() -> None:
    with pytest.raises(engine.NoTextError):
        engine.convert_document(make_docx_image_only(), ".docx", 50_000)


def test_pdf_chinese_text_and_table() -> None:
    result = engine.convert_document(make_chinese_pdf_with_table(), ".pdf", 50_000)
    md = result["markdown"]
    assert "张三" in md
    assert "研发部" in md
    assert result["warnings"] == []


def test_pdf_image_only_no_text() -> None:
    with pytest.raises(engine.NoTextError):
        engine.convert_document(make_pdf_image_only(), ".pdf", 50_000)


def test_pdf_encrypted() -> None:
    with pytest.raises(engine.EncryptedError):
        engine.convert_document(make_encrypted_pdf(), ".pdf", 50_000)


def test_pptx_with_picture_alt() -> None:
    # 只有一张带 alt 图的幻灯片 → 没有文字 → no_text
    with pytest.raises(engine.NoTextError):
        engine.convert_document(make_pptx_with_picture(), ".pptx", 50_000)


def test_pptx_image_only_no_text() -> None:
    with pytest.raises(engine.NoTextError):
        engine.convert_document(make_pptx_image_only(), ".pptx", 50_000)


def test_xlsx_basic_dates_and_numbers() -> None:
    result = engine.convert_document(make_xlsx_with_dates(), ".xlsx", 50_000)
    md = result["markdown"]
    assert "## 汇总" in md
    # 日期格式：纯日期 YYYY-MM-DD，带时间的留时间
    assert "2024-01-02" in md
    assert "2024-01-03 15:04:05" in md
    # 整数不带小数点
    assert "| 100 |" in md or "100 |" in md
    # 200.5 保留一位小数
    assert "200.5" in md
    # 12.0 这种整数浮点值不该有浮点噪声，也不该带 .0
    assert "12.0" not in md
    assert result["warnings"] == []


def test_xlsx_truncation_reports_real_total() -> None:
    data = make_xlsx_many_rows(200)
    result = engine.convert_document(data, ".xlsx", 120)  # 3 列 → 40 行（含表头）
    md = result["markdown"]
    assert "table_truncated" in result["warnings"]
    # 表尾就在截断位置后面不远，数得到：写出实际的总行数
    assert "此工作表共 200 行，只保留了前 39 行" in md


def test_xlsx_truncation_dimension_too_small() -> None:
    data = make_xlsx_bogus_dimension(200)
    result = engine.convert_document(data, ".xlsx", 120)
    md = result["markdown"]
    assert "table_truncated" in result["warnings"]
    # dimension 被篡改成 A1:C1（只有表头一行）：照样读到真实数据、数出真实的总行数
    assert "此工作表共 200 行，只保留了前 39 行" in md
    assert "| 0 | 0 | 0 |" in md
    assert "| 8 | 16 | 24 |" in md


def test_xlsx_truncation_dimension_too_large() -> None:
    # 整列套了格式时 dimension 会声明得比实际大很多，总行数不能照抄它
    data = make_xlsx_bogus_dimension(200, ref="A1:C100000")
    result = engine.convert_document(data, ".xlsx", 120)
    md = result["markdown"]
    assert "table_truncated" in result["warnings"]
    assert "此工作表共 200 行，只保留了前 39 行" in md
    assert "99999" not in md


def test_xlsx_truncation_too_long_to_count() -> None:
    # 比保留的行数多出 1000 行以上：不再往下数，只说原表更长
    data = make_xlsx_many_rows(2000)
    result = engine.convert_document(data, ".xlsx", 120)
    md = result["markdown"]
    assert "table_truncated" in result["warnings"]
    assert "此工作表只保留了前 39 行（原表更长）" in md
    assert "此工作表共" not in md


def test_xlsx_single_column_truncation() -> None:
    # 单列时行预算正好等于单元格预算：以前读不到触发截断的那一行，超出的行被静默丢掉
    data = make_xlsx_single_column(150)
    result = engine.convert_document(data, ".xlsx", 100)
    md = result["markdown"]
    assert "table_truncated" in result["warnings"]
    assert "此工作表共 150 行，只保留了前 99 行" in md
    assert "| 98 |" in md
    assert "| 99 |" not in md


def test_xlsx_single_column_exact_fit_not_truncated() -> None:
    # 表头 + 99 行正好 100 个单元格：装得下，不算截断
    data = make_xlsx_single_column(99)
    result = engine.convert_document(data, ".xlsx", 100)
    md = result["markdown"]
    assert result["warnings"] == []
    assert "只保留了前" not in md
    assert "| 98 |" in md


def test_xlsx_sheet_skip_after_budget() -> None:
    import io

    import openpyxl

    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Sheet1"
    ws1.append(["a", "b", "c"])
    for i in range(50):
        ws1.append([i, i * 2, i * 3])
    ws2 = wb.create_sheet("Sheet2")
    ws2.append(["x", "y"])
    for i in range(10):
        ws2.append([i, i * 2])
    buf = io.BytesIO()
    wb.save(buf)

    result = engine.convert_document(buf.getvalue(), ".xlsx", 120)
    md = result["markdown"]
    assert "## Sheet1" in md
    assert "## Sheet2" in md
    # Sheet2 被跳过：表头 + 「单元格预算已用完，此工作表已跳过。」
    assert "已跳过" in md


def test_csv_utf8() -> None:
    result = engine.convert_document(make_csv_utf8(), ".csv", 50_000)
    md = result["markdown"]
    assert "张三" in md
    assert "姓名" in md
    assert result["warnings"] == []


def test_csv_gbk_auto_decode() -> None:
    result = engine.convert_document(make_csv_gbk(), ".csv", 50_000)
    md = result["markdown"]
    assert "张三" in md
    assert "姓名" in md


def test_csv_truncation() -> None:
    data = make_csv_many_rows(500)
    result = engine.convert_document(data, ".csv", 150)
    md = result["markdown"]
    assert "table_truncated" in result["warnings"]
    assert "只保留了前" in md
    assert "此文件共 500 行" in md


def test_html_images_external_kept_internal_removed() -> None:
    result = engine.convert_document(make_html_with_images(), ".html", 50_000)
    md = result["markdown"]
    # 外链保持
    assert "https://example.com/pic.png" in md
    # 内嵌图 → 占位符
    assert "[图片：内嵌图]" in md
    assert "data:image/png;base64" not in md


def test_epub_chinese() -> None:
    result = engine.convert_document(make_epub(), ".epub", 50_000)
    md = result["markdown"]
    assert "测试书" in md
    assert "第一章" in md
    assert "中文内容测试" in md


def test_plain_zip_renamed_pdf_not_unpacked() -> None:
    # ZIP 改 .pdf：PdfConverter 会尝试解析并失败 → conversion_failed，而不是被别的
    # 转换器（比如 ZipConverter，我们没注册）解开
    with pytest.raises(engine.ConversionFailedError):
        engine.convert_document(make_plain_zip(), ".pdf", 50_000)


def test_bad_extension_no_converter() -> None:
    # .txt 没注册 PlainTextConverter → 不支持
    with pytest.raises(engine.UnsupportedFormatError):
        engine.convert_document(b"hello world", ".txt", 50_000)


def test_ole2_encrypted_detected_before_conversion() -> None:
    with pytest.raises(engine.EncryptedError):
        engine.convert_document(make_ole2_encrypted(), ".docx", 50_000)


def test_ole2_legacy_is_unsupported() -> None:
    with pytest.raises(engine.UnsupportedFormatError):
        engine.convert_document(make_ole2_legacy(), ".doc", 50_000)


def test_sniff_ole2_classification() -> None:
    assert sniff_ole2(b"hello world") is None
    assert sniff_ole2(make_ole2_legacy()) == "unsupported"
    assert sniff_ole2(make_ole2_encrypted()) == "encrypted"


def test_unreadable_glyphs_warning_threshold() -> None:
    # 9 个 (cid:...) 不够阈值
    md9 = "\n".join(f"(cid:{i})" for i in range(9))
    assert not postprocess.has_unreadable_glyphs(md9)
    # 10 个刚好触发
    md10 = "\n".join(f"(cid:{i})" for i in range(10))
    assert postprocess.has_unreadable_glyphs(md10)


def test_is_empty_of_text_strips_markdown_punctuation_and_images() -> None:
    assert postprocess.is_empty_of_text("")
    assert postprocess.is_empty_of_text("# ## ###")
    # 只有分隔线/竖线/空白，没有实际内容
    assert postprocess.is_empty_of_text("|  |  |  |\n|---|---|---|")
    # 表格里有数字/汉字才算"有文字"
    assert not postprocess.is_empty_of_text("| 姓名 | 分数 |\n|---|---|\n| 张三 | 100 |")
    assert not postprocess.is_empty_of_text("| a | b |\n|---|---|\n| 1 | 2 |")
    # 只有图片占位
    assert postprocess.is_empty_of_text("[图片]")
    assert postprocess.is_empty_of_text("[图片：foo]")
    # 幻灯片编号注释
    assert postprocess.is_empty_of_text("<!-- Slide number: 3 -->")
    assert postprocess.is_empty_of_text("  \n\t  \n")


def test_replace_images_md_and_html() -> None:
    text = (
        "![内嵌图](data:image/png;base64,abc) "
        "![外链](https://x.com/y.png) "
        '<img src="a.png" alt="内部图">'
    )
    result = postprocess.replace_images(text)
    assert "[图片：内嵌图]" in result
    assert "https://x.com/y.png" in result
    assert "[图片：内部图]" in result


def test_xlsx_pipe_escaping() -> None:
    import io

    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["a|b", "c"])
    ws.append(["d\ne", "f"])
    buf = io.BytesIO()
    wb.save(buf)
    result = engine.convert_document(buf.getvalue(), ".xlsx", 50_000)
    md = result["markdown"]
    # | 被转义，不会变成表格列分隔符
    assert "a\\|b" in md
    # 换行变成 <br>
    assert "d<br>e" in md
