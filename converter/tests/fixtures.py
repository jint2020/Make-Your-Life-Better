"""测试用的各种格式文件，全部在代码里生成，不提交二进制文件。

每个函数返回 bytes：构造函数尽量简单，够覆盖对应路径就行，不要追求真实文档的复杂度。
"""

from __future__ import annotations

import io
import zipfile
from datetime import datetime

import docx as _python_docx
import openpyxl
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.pdfencrypt import StandardEncryption
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen import canvas
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle

_OLE2_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def _small_png_bytes() -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (8, 8), color=(200, 30, 30)).save(buf, format="PNG")
    return buf.getvalue()


def make_docx_with_table_and_image() -> bytes:
    """一个含标题、段落、表格、一张内嵌图片（无 alt）的 docx。"""
    d = _python_docx.Document()
    d.add_heading("文档标题", level=1)
    d.add_paragraph("这是一段中文正文内容。")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "姓名"
    table.cell(0, 1).text = "分数"
    table.cell(1, 0).text = "张三"
    table.cell(1, 1).text = "100"
    d.add_picture(io.BytesIO(_small_png_bytes()))
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_docx_image_only() -> bytes:
    """只有一张图片的 docx，用来测 no_text。"""
    d = _python_docx.Document()
    d.add_picture(io.BytesIO(_small_png_bytes()))
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_chinese_pdf_with_table() -> bytes:
    """用 reportlab 的 STSong-Light CID 字体生成的中文 PDF，带一个表格。"""
    try:
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    except Exception:
        pass  # 重复注册会报错，忽略

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4)
    data = [["姓名", "部门", "分数"], ["张三", "研发部", "95"], ["李四", "市场部", "88"]]
    t = Table(data)
    t.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -1), "STSong-Light"),
                ("GRID", (0, 0), (-1, -1), 1, colors.black),
            ]
        )
    )
    doc.build([t])
    return buf.getvalue()


def make_pdf_image_only() -> bytes:
    """只有一张图片的 PDF，用来测 no_text（扫描件场景）。"""
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    from reportlab.lib.utils import ImageReader

    c.drawImage(ImageReader(io.BytesIO(_small_png_bytes())), 100, 700, width=100, height=100)
    c.save()
    return buf.getvalue()


def make_encrypted_pdf() -> bytes:
    """带密码的 PDF。"""
    buf = io.BytesIO()
    c = canvas.Canvas(buf, encrypt=StandardEncryption("secret-pass"))
    c.drawString(100, 750, "Encrypted content")
    c.save()
    return buf.getvalue()


def make_pptx_with_picture() -> bytes:
    """带一张有 alt 文本（descr）的图片的 PPTX。"""
    p = Presentation()
    slide = p.slides.add_slide(p.slide_layouts[6])
    pic = slide.shapes.add_picture(io.BytesIO(_small_png_bytes()), Inches(1), Inches(1))
    pic._element._nvXxPr.cNvPr.attrib["descr"] = "示意图"
    buf = io.BytesIO()
    p.save(buf)
    return buf.getvalue()


def make_pptx_image_only() -> bytes:
    """只有一张图片（没有 descr）的 PPTX，测 no_text。"""
    p = Presentation()
    slide = p.slides.add_slide(p.slide_layouts[6])
    slide.shapes.add_picture(io.BytesIO(_small_png_bytes()), Inches(1), Inches(1))
    buf = io.BytesIO()
    p.save(buf)
    return buf.getvalue()


def make_xlsx_with_dates(numbers: bool = True) -> bytes:
    """有日期、数字的 xlsx。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "汇总"
    ws.append(["姓名", "日期", "金额"])
    ws.append(["张三", datetime(2024, 1, 2), 100])
    ws.append(["李四", datetime(2024, 1, 3, 15, 4, 5), 200.5])
    if numbers:
        ws.append(["王五", datetime(2024, 1, 4), 12.0])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_xlsx_many_rows(num_rows: int) -> bytes:
    """很多行的 xlsx，用来测表截断。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "大表"
    ws.append(["a", "b", "c"])
    for i in range(num_rows):
        ws.append([i, i * 2, i * 3])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_xlsx_single_column(num_rows: int) -> bytes:
    """只有一列的 xlsx：单列时行预算正好等于单元格预算，专门覆盖截断的边界。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "单列"
    ws.append(["编号"])
    for i in range(num_rows):
        ws.append([i])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_xlsx_bogus_dimension(num_rows: int, ref: str = "A1:C1") -> bytes:
    """<dimension> 声明和实际行数对不上的 xlsx：默认标小；标大（比如 A1:C100000）也很常见，
    整列套了格式就会这样。验证我们不信声明值，读到、数出真实的行数。"""
    import re

    original = make_xlsx_many_rows(num_rows)
    src = zipfile.ZipFile(io.BytesIO(original))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for item in src.infolist():
            content = src.read(item.filename)
            if item.filename.startswith("xl/worksheets/sheet"):
                content = re.sub(
                    rb'<dimension ref="[^"]*"/>',
                    f'<dimension ref="{ref}"/>'.encode(),
                    content,
                )
            z.writestr(item, content)
    return out.getvalue()


def make_csv_utf8() -> bytes:
    return "姓名,分数\n张三,100\n李四,95\n".encode()


def make_csv_gbk() -> bytes:
    return "姓名,分数\n张三,100\n李四,95\n".encode("gbk")


def make_csv_many_rows(num_rows: int) -> bytes:
    lines = ["a,b,c"]
    for i in range(num_rows):
        lines.append(f"{i},{i * 2},{i * 3}")
    return "\n".join(lines).encode("utf-8")


def make_html_with_images() -> bytes:
    import base64

    png_b64 = base64.b64encode(_small_png_bytes()).decode()
    return f"""<html><body>
<h1>标题</h1>
<p>正文文字</p>
<img src="data:image/png;base64,{png_b64}" alt="内嵌图">
<img src="https://example.com/pic.png" alt="外链图">
</body></html>""".encode()


def make_epub() -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("mimetype", "application/epub+zip", zipfile.ZIP_STORED)
        container_xml = (
            '<?xml version="1.0"?>\n'
            '<container version="1.0" '
            'xmlns="urn:oasis:names:tc:opendocument:xmlns:container">\n'
            '  <rootfiles><rootfile full-path="OEBPS/content.opf" '
            'media-type="application/oebps-package+xml"/></rootfiles>\n'
            "</container>"
        )
        z.writestr("META-INF/container.xml", container_xml)
        z.writestr(
            "OEBPS/content.opf",
            """<?xml version="1.0"?>
<package xmlns="http://www.idpf.org/2007/opf" version="2.0" unique-identifier="BookId">
  <metadata xmlns:dc="http://purl.org/dc/elements/1.1/">
    <dc:title>测试书</dc:title>
    <dc:language>zh</dc:language>
  </metadata>
  <manifest>
    <item id="c1" href="chapter1.xhtml" media-type="application/xhtml+xml"/>
  </manifest>
  <spine><itemref idref="c1"/></spine>
</package>""",
        )
        z.writestr(
            "OEBPS/chapter1.xhtml",
            """<?xml version="1.0" encoding="UTF-8"?>
<html xmlns="http://www.w3.org/1999/xhtml"><body><h1>第一章</h1><p>中文内容测试。</p></body></html>""",
        )
    return buf.getvalue()


def make_plain_zip() -> bytes:
    """一个普通的 zip（非 Office），改扩展名后应该被拒绝，不能被解开。"""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("hello.txt", "just a plain zip, not an office document" * 5)
    return buf.getvalue()


def make_ole2_encrypted() -> bytes:
    """模拟一个加密的 OOXML：OLE2 magic + EncryptedPackage UTF-16LE 流名。"""
    return _OLE2_MAGIC + b"\x00" * 100 + "EncryptedPackage".encode("utf-16-le") + b"\x00" * 100


def make_ole2_legacy() -> bytes:
    """模拟一个旧版二进制 Office 文件（OLE2 但没有 EncryptedPackage 流）。"""
    return _OLE2_MAGIC + b"\x00" * 500
